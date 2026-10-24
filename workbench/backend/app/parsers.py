"""Bounded file parsers. No uploaded macros, formulas or executables are run."""
from __future__ import annotations
import csv, io, math, re, zipfile
from datetime import datetime, date
from pathlib import Path
from decimal import Decimal
from openpyxl import load_workbook
from .domain import DomainError, MAX_BYTES, MAX_ROWS, MAX_COLUMNS


def cell_text(value: object) -> str:
    if value is None: return ''
    if isinstance(value, bool): return 'TRUE' if value else 'FALSE'
    if isinstance(value,(date,datetime)): return value.date().isoformat() if isinstance(value,datetime) else value.isoformat()
    if isinstance(value,float):
        if not math.isfinite(value): raise DomainError('nonfinite_cell','单元格含非有限数值。')
        return format(Decimal(str(value)), 'f').rstrip('0').rstrip('.') if '.' in format(Decimal(str(value)),'f') else str(int(value))
    return str(value)


def parse_file(data: bytes, filename: str, *, sheet: str | None = None, header_row: int = 1) -> dict:
    if not data or len(data)>MAX_BYTES: raise DomainError('file_size','文件为空或超过8 MiB限制。')
    suffix=Path(filename).suffix.lower()
    warnings=[]; numeric_cells=[]
    if data[:8]==bytes.fromhex('d0cf11e0a1b11ae1'):
        if suffix!='.xls': raise DomainError('format_mismatch','真实格式是二进制XLS，后缀不一致。')
        # Preflight BIFF records, including formulas, before xlrd can expose stale caches.
        from .xls_safe import read_xls
        try: parsed=read_xls(data,sheet)
        except DomainError: raise
        except Exception: raise DomainError('invalid_xls','XLS结构损坏或含不支持的特性。')
        raw=parsed['rows']; name=parsed['sheet']; names=parsed['sheets']; engine=parsed['engine']; numeric_cells=parsed['numeric_cells']
    elif data[:2]==b'PK':
        if suffix!='.xlsx': raise DomainError('format_mismatch','仅支持真实XLSX，不接受宏文件或伪装后缀。')
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                entries=z.infolist(); names_zip=z.namelist()
                if len(entries)>1000 or sum(x.file_size for x in entries)>24*1024*1024:
                    raise DomainError('zip_limit','压缩包展开体积或成员数超过限制。')
                if any(x.file_size>MAX_BYTES or (x.compress_size and x.file_size/x.compress_size>250) for x in entries):
                    raise DomainError('zip_limit','压缩成员过大或压缩比例异常。')
                if len(set(names_zip))!=len(names_zip): raise DomainError('zip_duplicate','压缩包有重复成员。')
                if any('vba' in x.lower() or 'externallink' in x.lower() or 'embedding' in x.lower() or 'activex' in x.lower() for x in names_zip):
                    raise DomainError('active_content','不接受宏、嵌入对象、ActiveX或外部链接。')
                if '[Content_Types].xml' not in names_zip or 'xl/workbook.xml' not in names_zip: raise DomainError('invalid_xlsx','不是有效XLSX。')
                types=z.read('[Content_Types].xml').lower()
                if b'macroenabled' in types: raise DomainError('active_content','不接受宏工作簿。')
            wb=load_workbook(io.BytesIO(data),read_only=False,data_only=False,keep_links=False)
            names=wb.sheetnames
            if len(names)>1 and sheet is None: raise DomainError('choose_sheet','工作簿含多个工作表，请明确选择。')
            name=sheet or names[0]
            if name not in names: raise DomainError('sheet_not_found','指定工作表不存在。')
            ws=wb[name]
            if ws.sheet_state!='visible': raise DomainError('hidden_sheet','不接受隐藏工作表。')
            if ws.max_row>MAX_ROWS+30 or ws.max_column>MAX_COLUMNS: raise DomainError('row_limit','行数或列数超过原型限制。')
            if ws.merged_cells.ranges: raise DomainError('merged_cells','含合并单元格：请先按经确认的模板展开；本版不自动填充。')
            raw=[]
            for rn, row in enumerate(ws.iter_rows(),1):
                vals=[]
                if ws.row_dimensions[rn].hidden: raise DomainError('hidden_row','含隐藏行，不能默认为完整导出。')
                for cn,c in enumerate(row,1):
                    if c.data_type=='f': raise DomainError('formula_cell','含公式；本版要求静态导出值，不使用可能过期的缓存。')
                    if c.data_type=='e': raise DomainError('error_cell','工作簿含错误单元格。')
                    if isinstance(c.value,(int,float)) and not isinstance(c.value,bool):
                        numeric_cells.append([rn,cn])
                        if re.fullmatch(r'0{2,}',c.number_format or '') and float(c.value).is_integer():
                            vals.append(str(int(c.value)).zfill(len(c.number_format))); continue
                    vals.append(cell_text(c.value))
                raw.append(vals)
            wb.close(); engine='openpyxl'
        except DomainError: raise
        except Exception: raise DomainError('invalid_xlsx','XLSX损坏、加密或无法解析。')
    else:
        if suffix!='.csv': raise DomainError('format_mismatch','文件签名不匹配；不把文本伪装成XLS。')
        if b'\x00' in data: raise DomainError('binary_csv','CSV含NUL；仅接受UTF-8或GB18030文本。')
        try: text=data.decode('utf-8-sig'); engine='csv-utf8'
        except UnicodeDecodeError:
            try: text=data.decode('gb18030'); engine='csv-gb18030'
            except UnicodeDecodeError: raise DomainError('encoding','CSV编码无法识别。')
        if text.lstrip().lower().startswith(('<html','<!doctype','<?xml')): raise DomainError('html_export','这是HTML/XML报表，不是CSV/XLS。')
        if len(text)>MAX_BYTES: raise DomainError('text_limit','文本过长。')
        try:
            reader=csv.reader(io.StringIO(text),strict=True); raw=[]
            for row in reader:
                if len(row)>MAX_COLUMNS or len(raw)>MAX_ROWS+30: raise DomainError('row_limit','行数或列数超过限制。')
                raw.append(row)
        except csv.Error: raise DomainError('invalid_csv','CSV引号或字段结构错误。')
        name='CSV'; names=['CSV']
    if not 1<=header_row<=30: raise DomainError('header_row','表头行须在1到30之间。')
    if not raw or header_row>len(raw): raise DomainError('empty_sheet','没有可读表头。')
    headers=[str(x).strip() for x in raw[header_row-1]]
    # Trim trailing wholly empty columns, not intermediate headers or data.
    while headers and not headers[-1] and all(len(r)<len(headers) or not str(r[len(headers)-1]).strip() for r in raw[header_row:]): headers.pop()
    if not headers or any(not h for h in headers): raise DomainError('blank_header','表头存在空字段，请明确整理模板。')
    if len(set(headers))!=len(headers): raise DomainError('duplicate_header','表头字段重复，无法可靠映射。')
    records=[]; blank=[]
    for rn, values in enumerate(raw[header_row:],header_row+1):
        if all(not str(x).strip() for x in values): blank.append(rn); continue
        if len(values)>len(headers): raise DomainError('column_mismatch',f'第{rn}行列数超出表头。')
        values=list(values)+['']*(len(headers)-len(values))
        records.append({'row':rn,'values':dict(zip(headers,values))})
    if len(records)>MAX_ROWS: raise DomainError('row_limit','单批次最多5000条记录。')
    if not records: raise DomainError('empty_data','没有数据行。')
    if blank: warnings.append({'code':'blank_rows','rows':blank,'message':'空白行已显式记录，未计入业务行数。'})
    if numeric_cells: warnings.append({'code':'numeric_cells','message':'数值型标识可能已丢失前导零；仅恢复明确的全零显示格式，请核对预览。','count':len(numeric_cells)})
    return {'format':suffix[1:],'engine':engine,'sheet':name,'sheets':names,'headers':headers,'rows':records,'warnings':warnings,'blank_rows':blank}
